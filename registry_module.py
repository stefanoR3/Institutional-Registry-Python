import pymysql #used to send MySql network requests using python commands
import openpyxl as xl
import openpyxl.styles as xs #not included in the init file of the openpyxl package
from openpyxl.utils import get_column_letter
import datetime
import pandas_module as pd_mod
import pandas as pd

def get_connection() -> pymysql.Connection:
    connection = pymysql.connect(
        host='localhost',
        user='root',
        password='',
        database='regtest',
        port = 3306,
        charset='utf8mb4',
        cursorclass=pymysql.cursors.DictCursor #convert tuples into dictionaries
    )
    return connection

def fetch_registry_records(start_date:datetime.date = None, end_date:datetime.date = None) -> tuple[tuple]:
    #fetch rows trough a data range
    connection = get_connection()
    with connection.cursor() as cursor: #study cursors!
        #we prevent sql injection by not letting further instruction append
        # ''' or """ used for multi-line strings
        if start_date==None and end_date==None:
            stmt = """SELECT * FROM registru"""
            cursor.execute(stmt)
        elif start_date==None and end_date != None:
            stmt = """SELECT iddoc, nrinreg, dtinreg, subnrinreg, subdtinreg, inties, modintrare, tipintrare, termen_zile, termen_data, emitent, emailemitent, nrinregemit, dtinregemit, destinatar, emaildestinatar, compartiment_ies, continut_int,
              continut_ies, repartdanu,repartizat, dtrepar, tipiesire, modexp, solras_danu, solras_data, nrinregdest, dtinregdest, nrinreg_cnx, dtinreg_cnx, inddos, mot_anul, stare, utilizator, last_update FROM registru WHERE dtinreg<%s""" #injection?
            cursor.execute(stmt,(end_date))
        elif start_date != None and end_date==None:
            stmt = """SELECT iddoc, nrinreg, dtinreg, subnrinreg, subdtinreg, inties, modintrare, tipintrare, termen_zile, termen_data, emitent, emailemitent, nrinregemit, dtinregemit, destinatar, emaildestinatar, compartiment_ies, continut_int,
              continut_ies, repartdanu,repartizat, dtrepar, tipiesire, modexp, solras_danu, solras_data, nrinregdest, dtinregdest, nrinreg_cnx, dtinreg_cnx, inddos, mot_anul, stare, utilizator, last_update FROM registru WHERE dtinreg>%s"""
            cursor.execute(stmt,(start_date))
        else:
            stmt = """SELECT iddoc, nrinreg, dtinreg, subnrinreg, subdtinreg, inties, modintrare, tipintrare, termen_zile, termen_data, emitent, emailemitent, nrinregemit, dtinregemit, destinatar, emaildestinatar, compartiment_ies, continut_int,
              continut_ies, repartdanu,repartizat, dtrepar, tipiesire, modexp, solras_danu, solras_data, nrinregdest, dtinregdest, nrinreg_cnx, dtinreg_cnx, inddos, mot_anul, stare, utilizator, last_update FROM registru WHERE dtinreg BETWEEN %s AND %s"""
            cursor.execute(stmt, (start_date,end_date))

        records = cursor.fetchall()
        #for record in records:
            #print(f"ID:{record['iddoc']} | Nr:{record['nrinreg']} | Date:{record['dtinreg']} |  Emitent:{record['emitent']} | continut_int:{record['continut_int']}")
            #print(type('dtinreg'))

        return records
    
def registry_export(df:pd.DataFrame):
    workbook = xl.Workbook() #xcel container for all sheets(folder)
    worksheet = workbook.active #single tab/page inside the workbook
    worksheet.title  = "REGISTRUL DE INTRĂRI-IEȘIRI"
    worksheet.page_setup.showGridLines = True
    worksheet.views.sheetView[0].showGridLines = True #both on screen/page
    headers = list(df.columns)
    worksheet.append(headers) #write headers on row 1
    for row in df.itertuples(index=False): #at every append it creates a new row
        worksheet.append(list(row)) #itertuples for namedTuples, iterows for series

    #shifts headers to row 3
    worksheet.insert_rows(1, amount=2)
    worksheet.merge_cells(start_row=1, start_column=1, end_column=len(headers), end_row=2)
    min_data = df['din data'].min()
    max_data = df['din data'].max()

    title_cell = worksheet.cell(column=1, row=1)
    title_cell.value = f"REGISTRUL DE INTRĂRI-IEȘIRI \n DIN PERIOADA {min_data}-{max_data} "
    title_cell.font = xs.Font(name='Calibri', size=11, bold=True)
    title_cell.alignment = xs.Alignment(horizontal='center', vertical='center', wrap_text=True)
    worksheet.row_dimensions[1].height = 20
    worksheet.row_dimensions[2].height = 20 #?

    #format header row ([3])
    header_font = xs.Font(name='Calibri', size=10, bold=True)
    header_alignment = xs.Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_side = xs.Side(style='thin', color='A0A0A0')
    grid_border = xs.Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for cell in worksheet[3]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = grid_border

    worksheet.row_dimensions[3].height = 32

    left_align_wrap = xs.Alignment(horizontal='left', vertical='center', wrap_text=True)
    center_align_wrap = xs.Alignment(horizontal='center', vertical='center', wrap_text=True)

    #applying style on row 4 onwords (the row element not modified in enum)
    for row in worksheet.iter_rows(min_row=4, max_row=worksheet.max_row, min_col=1, max_col=len(headers)):
        for col_index, cell in enumerate(row, start=0): #enum merge in dreapta (1,1), (1,2) etc.
            cell.border = grid_border
            cell.font = xs.Font(name='Calibri', size=10)
            need_extra_height = False
            row_num = row[0].row #first cell row index

            #center align short metadata
            header_name = str(headers[col_index]).lower()
            if any(k in header_name for k in ['număr','din data','data expedierii','indicator nomenelator']):
                cell.alignment = center_align_wrap
            else:
                cell.alignment = left_align_wrap

            if cell.value:
                text = str(cell.value)
                if '\n' in text or len(text)>30:
                    need_extra_height = True

            if need_extra_height:
                worksheet.row_dimensions[row_num].height = 28


    #column widths
    for col_idx in range(1, len(headers)+1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in worksheet[col_letter]:
            if cell.value:
                lines = str(cell.value).split('\n')
                max_len = max(max_len, max(len(line) for line in lines))

        #clamp
        worksheet.column_dimensions[col_letter].width = min(max(max_len + 3,10), 35)
        
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.4
    worksheet.page_margins.bottom = 0.4

    workbook.save('registru_export.xlsx')
    
    

